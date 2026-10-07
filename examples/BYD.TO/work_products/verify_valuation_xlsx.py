"""
Recalculate work_products/valuation.xlsx from its formulas (no Excel needed - uses the `formulas` package) and check
every key output against the python scripts it mirrors:

    owners_earnings_roic.py  -> sheet 1 section 2 (historical OE) and 2b (CFO reconciliation)
    forecast_earnings.py     -> sheet 1 section 4 (forecast workings, all four cases)
    terminal_value.py        -> sheet 2
    valuation.py             -> sheet 3 (SOTP, perpetuity cross-check, sensitivities, bridge alternatives)

    uv run python work_products/verify_valuation_xlsx.py

Tolerance: the workbook links the forecast to UNROUNDED filing figures (e.g. FY2025 sales 3,142.794 where
forecast_earnings.py types 3,142.8; H1-2026 rent 107.941 where valuation.py types 107.9), so small differences
are expected. A check passes if |diff| <= max(0.05, 0.1% of the value).
"""
from __future__ import annotations

import contextlib
import io
import os
import sys

import formulas
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_valuation_xlsx as bx  # noqa: E402
import forecast_earnings as fe  # noqa: E402
import terminal_value as tv  # noqa: E402
import valuation as val  # noqa: E402

with contextlib.redirect_stdout(io.StringIO()):
    import owners_earnings_roic as oer  # noqa: E402  (prints on import)

earn, tvs, vs = bx.main()
BOOK = os.path.basename(bx.OUT_PATH).upper()

print("recalculating workbook with `formulas` ...")
model = formulas.ExcelModel().loads(bx.OUT_PATH).finish()
solution = model.calculate()
VALUES = {}
for key, v in solution.items():
    try:
        VALUES[key.upper()] = v.value[0, 0]
    except Exception:  # ranges
        pass


def cell(sb: bx.SheetBuilder, name: str, col: str):
    key = f"'[{BOOK}]{sb.title.upper()}'!{col}{sb.rows[name]}"
    return VALUES[key]


failures, n_checks = [], 0


def check(label: str, got, expected, abs_tol=0.05, rel_tol=1e-3):
    global n_checks
    n_checks += 1
    ok = isinstance(got, (int, float)) and abs(got - expected) <= max(abs_tol, rel_tol * abs(expected))
    if not ok:
        failures.append(f"{label}: workbook {got!r} vs python {expected!r}")


# ---------------------------------------------------------------- 0. no error values, no '=' labels
errors = [k for k, v in VALUES.items() if isinstance(v, str) and v.startswith("#")]
wb = openpyxl.load_workbook(bx.OUT_PATH)
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.column_letter in ("A", "L") and isinstance(c.value, str) and c.value.startswith("="):
                errors.append(f"{ws.title}!{c.coordinate} label starts with '='")
print(f"error cells / formula-like labels: {len(errors)}")
failures += errors

# ---------------------------------------------------------------- 1. historical OE (sheet 1 section 2)
cols = dict(zip(oer.PERIODS, bx.HIST_COLS))
for p, col in cols.items():
    row = oer.rows_oe[p]
    check(f"hist {p} sales", cell(earn, "h_sales", col), row["Sales"] / 1000)
    check(f"hist {p} adj EBITDA", cell(earn, "h_ebitda", col), row["Adjusted EBITDA"] / 1000)
    check(f"hist {p} rent", cell(earn, "h_lease", col), row["(-) Cash lease payments"] / 1000)
    check(f"hist {p} pre-tax OE", cell(earn, "h_pretax", col), row["Pre-tax Owners' Earnings"] / 1000)
    check(f"hist {p} OE", cell(earn, "h_oe", col), row["Owners' Earnings"] / 1000)
    check(f"hist {p} CFO recon OE", cell(earn, "r_oe", col), oer.recon[p]["= Owners' Earnings (normalized)"] / 1000)
    check(f"hist {p} CFO recon check", cell(earn, "r_check", col), 0.0, abs_tol=1e-6)
    check(f"hist {p} OE at cash tax", cell(earn, "h_oe_cashtax", col), oer.owners_earnings_if_cash_tax_used(p) / 1000)

# ---------------------------------------------------------------- 2. forecast (sheet 1 section 4)
for case in bx.CASES:
    years = fe.forecast(bx.SCEN_DEF[case])
    p = case + "_"
    for y, col in zip(years, bx.YEAR_COLS):
        tag = f"{case} {y.year}"
        check(f"{tag} SSS", cell(earn, p + "sss", col), y.sss, abs_tol=1e-9)
        check(f"{tag} existing sales", cell(earn, p + "ex_sales", col), y.existing.sales)
        check(f"{tag} existing EBITDA", cell(earn, p + "ex_ebitda", col), y.existing.ebitda)
        check(f"{tag} rent", cell(earn, p + "rent", col), y.existing.rent)
        check(f"{tag} new-loc EAR", cell(earn, p + "new_ear", col), y.new.ebitda_after_rent)
        check(f"{tag} new-loc sales", cell(earn, p + "new_sales", col), y.new.sales)
        check(f"{tag} run-rate EAR", cell(earn, p + "new_rr_ear", col), y.new.run_rate_ebitda_after_rent)
        check(f"{tag} growth capex", cell(earn, p + "gcapex", col), y.new.growth_capex)
        check(f"{tag} total sales", cell(earn, p + "tot_sales", col), y.sales)
        check(f"{tag} total EBITDA memo", cell(earn, p + "tot_ebitda", col), y.ebitda_memo)
        check(f"{tag} EBITDA after rent", cell(earn, p + "ear", col), y.ebitda_after_rent)
        check(f"{tag} OE", cell(earn, p + "oe", col), y.owners_earnings)
        check(f"{tag} OE existing", cell(earn, p + "oe_ex", col), y.oe_existing)
        check(f"{tag} OE new", cell(earn, p + "oe_new", col), y.oe_new)
        check(f"{tag} OE new run-rate", cell(earn, p + "oe_new_rr", col), y.oe_new_run_rate)
        check(f"{tag} one-offs", cell(earn, p + "oneoffs", col), y.one_offs)
        check(f"{tag} distributable", cell(earn, p + "dist", col), y.distributable)
        check(f"{tag} net debt", cell(earn, p + "netdebt", col), y.net_debt)
        check(f"{tag} summary OE", cell(earn, "S_" + p + "oe", col), y.owners_earnings)
        s_ear = cell(earn, "S_" + p + "ebitda", col) + cell(earn, "S_" + p + "rent", col)
        check(f"{tag} summary bridge EBITDA - rent = EAR", s_ear, y.ebitda_after_rent)
res = {n: fe.forecast(s) for n, s in zip(("Bear", "Base", "Bull"), fe.SCENARIOS)}
for i, col in enumerate(bx.YEAR_COLS):
    check(f"weighted OE {fe.YEARS[i]}", cell(earn, "P_oe", col), fe.probability_weighted(res, lambda y: y.owners_earnings)[i])

# ---------------------------------------------------------------- 3. terminal value (sheet 2)
for case in bx.VAL_CASES:
    col, name = bx.VAL_COL[case], bx.TV_KEY[case]
    t = tv.terminal_value(name)
    check(f"TV {case} ramp gap", cell(tvs, "gap", col), t["gap"])
    check(f"TV {case} total", cell(tvs, "tv", col), t["total"])
    check(f"TV {case} implied r", cell(tvs, "vd_implied_r", col),
          tv.implied_discount_rate(t["total"] / t["oe"], tv.STEADY[name]), abs_tol=1e-4)
    check(f"TV {case} fundamental multiple", cell(tvs, "vd_mult", col), tv.fundamental_oe_multiple(tv.STEADY[name]))
tvd = {n: tv.terminal_value(n) for n in tv.TERMINAL}
check("TV weighted total", cell(tvs, "tv", bx.WTD_COL), tv.probability_weighted(tvd, "total"))
base_r, ng_r = tv.RESULTS["Base"], tv.RESULTS[tv.NO_GROWTH]
d_dist = {b.year: b.distributable - n.distributable for b, n in zip(base_r, ng_r)}
check("growth programme net value", cell(tvs, "gp_net", "H"),
      tvd["Base"]["total"] - tvd[tv.NO_GROWTH]["total"] + tv.future_value_2031(d_dist))
jd = tv.JoeHudsonsDeal()
check("JH implied synergies", cell(tvs, "jh_syn", "B"), jd.implied_synergies)

# ---------------------------------------------------------------- 4. valuation (sheet 3)
r = val.SETTINGS.discount_rate
exit_vals = val.scenario_values(r, "exit")
perp_vals = val.scenario_values(r, "perpetuity")
for case in bx.VAL_CASES:
    col, name = bx.VAL_COL[case], bx.TV_KEY[case]
    check(f"val {case} stub", cell(vs, "stub", col), val.stub_2026_distributable(val.FORECASTS[name][0])["distributable"])
    check(f"val {case} PV CF", cell(vs, "v_pvcf", col), exit_vals[name]["pv_cash_flows"])
    check(f"val {case} EV", cell(vs, "v_ev", col), exit_vals[name]["ev"])
    check(f"val {case} equity", cell(vs, "v_eq", col), exit_vals[name]["equity"])
    check(f"val {case} per share", cell(vs, "v_ps", col), exit_vals[name]["per_share"], abs_tol=0.02)
    check(f"val {case} perpetuity per share", cell(vs, "p_ps", col), perp_vals[name]["per_share"], abs_tol=0.02)
wtd = val.weighted({n: exit_vals[n]["per_share"] for n in val.WEIGHTS})
check("HEADLINE weighted US$/share", cell(vs, "v_ps", bx.WTD_COL), wtd, abs_tol=0.02)
check("HEADLINE weighted C$/share", cell(vs, "v_ps_cad", bx.WTD_COL), wtd * val.FX.cad_per_usd, abs_tol=0.03)
check("weighted perpetuity US$/share", cell(vs, "p_ps", bx.WTD_COL),
      val.weighted({n: perp_vals[n]["per_share"] for n in val.WEIGHTS}), abs_tol=0.02)
check("standalone assets", cell(vs, "b_sa", "B"), val.standalone_assets(r)["total"])
check("debt total", cell(vs, "b_debt", "B"), val.debt_items()["total"])
check("diluted shares (m)", cell(vs, "b_shares", "B"), val.SHARES.diluted_shares / 1e6, abs_tol=1e-6)
check("JH own-formula PV", cell(vs, "jt_pv", "B"), val.pv_jh_tax_shield(r))
check("memo r1 weights per share", cell(vs, "v_ps_r1", bx.WTD_COL),
      sum(val.PRIOR_WEIGHTS[n] * exit_vals[n]["per_share"] for n in val.PRIOR_WEIGHTS), abs_tol=0.02)
for i, rate in enumerate((0.08, 0.09, 0.10, 0.11, 0.12)):
    ex, pp = val.scenario_values(rate, "exit"), val.scenario_values(rate, "perpetuity")
    check(f"sens exit weighted @ {rate:.0%}", cell(vs, f"sr_exit{i}", "G"),
          val.weighted({n: ex[n]["per_share"] for n in val.WEIGHTS}), abs_tol=0.02)
    check(f"sens perp weighted @ {rate:.0%}", cell(vs, f"sr_perp{i}", "G"),
          val.weighted({n: pp[n]["per_share"] for n in val.WEIGHTS}), abs_tol=0.02)
for i, shift in enumerate((-1.5, -1.0, -0.5, 0.0, 0.5, 1.0)):
    v = val.scenario_values(r, "exit", shift)
    check(f"sens multiple shift {shift:+.1f}", cell(vs, f"sm{i}", "G"),
          val.weighted({n: v[n]["per_share"] for n in val.WEIGHTS}), abs_tol=0.02)
w_eq = val.weighted({n: exit_vals[n]["equity"] for n in val.WEIGHTS})
check("alt: deduct net DTL", cell(vs, "alt1", "C"),
      val.per_share(w_eq - (val.DEBT.deferred_tax_liability_memo - val.DEBT.deferred_tax_asset_memo)), abs_tol=0.02)
check("alt: drop all shields", cell(vs, "alt2", "C"), val.per_share(w_eq - val.tax_shield_total(r)), abs_tol=0.02)
check("alt: leases deducted", cell(vs, "alt6", "C"), val.per_share(w_eq - val.DEBT.lease_liabilities_memo), abs_tol=0.02)

# ---------------------------------------------------------------- 5. live-link test: change inputs, value moves
# Edit input cells in a temporary copy, recalculate, and compare with python run on the same inputs.
import tempfile  # noqa: E402
from dataclasses import replace as dc_replace  # noqa: E402


def recalc_with(edits: dict) -> dict:
    wb2 = openpyxl.load_workbook(bx.OUT_PATH)
    for (sheet, addr), value in edits.items():
        wb2[sheet][addr] = value
    tmp = os.path.join(tempfile.mkdtemp(), os.path.basename(bx.OUT_PATH))
    wb2.save(tmp)
    sol = formulas.ExcelModel().loads(tmp).finish().calculate()
    return {k.upper(): v.value[0, 0] for k, v in sol.items() if hasattr(v, "value") and v.value.size == 1}


ps_key = f"'[{BOOK}]{vs.title.upper()}'!{bx.WTD_COL}{vs.rows['v_ps']}"
sol9 = recalc_with({(vs.title, f"B{vs.rows['r']}"): 0.09})
ex9 = val.scenario_values(0.09, "exit")
check("live link: discount rate 9% -> weighted per share", sol9[ps_key],
      val.weighted({n: ex9[n]["per_share"] for n in val.WEIGHTS}), abs_tol=0.02)

# maintenance capex 2.0% (sheet 1 input 1.2) flows to historical OE and to the Base forecast (linked), and on to value
sol_m = recalc_with({(earn.title, f"B{earn.rows['a_maint']}"): 0.020})
oe_key = f"'[{BOOK}]{earn.title.upper()}'!D{earn.rows['h_oe']}"
check("live link: maint capex 2.0% -> FY2025 OE", sol_m[oe_key], oer.sens[0.020]["FY2025"][0] / 1000)
base_m = fe.forecast(dc_replace(fe.BASE, maintenance_capex_pct_sales=0.020))
k31 = f"'[{BOOK}]{earn.title.upper()}'!K{earn.rows['Base_oe']}"
check("live link: maint capex 2.0% -> Base FY2031E OE", sol_m[k31], base_m[-1].owners_earnings)

print(f"\n{n_checks} checks, {len(failures)} failures")
for f in failures[:60]:
    print("  FAIL", f)
print(f"\nHeadline (workbook): US${cell(vs, 'v_ps', bx.WTD_COL):.2f} per share "
      f"(C${cell(vs, 'v_ps_cad', bx.WTD_COL):.2f}); python US${wtd:.2f}")
print("  by case: " + ", ".join(f"{c} US${cell(vs, 'v_ps', bx.VAL_COL[c]):.2f}" for c in bx.VAL_CASES))
print(f"  perpetuity cross-check weighted: US${cell(vs, 'p_ps', bx.WTD_COL):.2f}")
sys.exit(1 if failures else 0)
