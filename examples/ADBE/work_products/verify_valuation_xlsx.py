"""
Adobe Inc. (ADBE) — verifies work_products/valuation.xlsx.

The workbook contains only formulas outside the Inputs sheet, so "does it tie?"
cannot be answered by reading the file. This script RECALCULATES every formula
with the `formulas` engine (a pure-python Excel evaluator) and then:

  1. asserts every internal consistency check evaluates to ~0
       - the driver-built income statement reproduces reported EBIT
       - the three customer groups sum to total subscription revenue
       - scenario probabilities sum to 100%
       - weighting-then-discounting equals discounting-then-weighting
  2. compares every "Model vs Published" check row against the figure published
     in the underlying work product, and reports the difference.

A non-zero residual is acceptable ONLY where it is explained: the published
scripts round scenario owners' earnings to the nearest $M before weighting,
whereas the workbook carries full precision throughout.

Run:  uv run python work_products/verify_valuation_xlsx.py
"""

import sys
import warnings

import formulas
import openpyxl

warnings.filterwarnings("ignore")

XLSX = "work_products/valuation.xlsx"
TOL_ABS = 1.0        # $M / $ per share: residual tolerated from rounding
TOL_REL = 0.001      # 0.1% relative tolerance on large aggregates


def load():
    wb = openpyxl.load_workbook(XLSX)
    xl = formulas.ExcelModel().loads(XLSX).finish()
    sol = xl.calculate()
    vals = {}
    for k, v in sol.items():
        if "]" not in k or "'!" not in k:
            continue
        sheet = k.split("]", 1)[1].split("'!", 1)[0].upper()
        ref = k.rsplit("'!", 1)[1]
        try:
            vals[(sheet, ref)] = v.value[0, 0]
        except Exception:
            pass
    return wb, vals


def find_rows(ws, needle, col=1):
    """Row numbers whose column-A label contains `needle`."""
    out = []
    for row in range(1, ws.max_row + 1):
        v = ws.cell(row=row, column=col).value
        if isinstance(v, str) and needle.lower() in v.lower():
            out.append(row)
    return out


def main():
    wb, vals = load()
    sh_in, sh_oe, sh_tv = wb.sheetnames
    fails, warns = [], []

    def get(sheet, ref):
        return vals.get((sheet.upper(), ref))

    print("=" * 96)
    print("PART 1 — INTERNAL CONSISTENCY CHECKS (all must be ~0)")
    print("=" * 96)
    ws = wb[sh_oe]
    zero_checks = []
    for row in find_rows(ws, "must be 0"):
        label = ws.cell(row=row, column=1).value.strip()
        for col in "BCDE":
            if ws.cell(row=row, column=openpyxl.utils.column_index_from_string(col)
                       ).value is None:
                continue
            zero_checks.append((sh_oe, label, f"{col}{row}"))
    for row in find_rows(wb[sh_tv], "must be 0"):
        label = wb[sh_tv].cell(row=row, column=1).value.strip()
        zero_checks.append((sh_tv, label, f"F{row}"))
    # probabilities sum to 100%
    for row in find_rows(wb[sh_in], "probabilities sum"):
        zero_checks.append((sh_in, "probabilities sum to 100% (shown as 1.0)",
                            f"B{row}"))
    # EBITA reconciliation gap (documented, non-zero)
    gap_rows = find_rows(ws, "less the §1 bridge EBITA")

    for sheet, label, ref in zero_checks:
        v = get(sheet, ref)
        target = 1.0 if "probabilities" in label else 0.0
        ok = v is not None and abs(v - target) < 1e-6
        print(f"  [{'OK ' if ok else 'FAIL'}]  {sheet:<26} {ref:<6} "
              f"{label[:58]:<58} = {v}")
        if not ok:
            fails.append((label, ref, v, target))

    for row in gap_rows:
        v = get(sh_oe, f"D{row}")
        print(f"  [NOTE]  {sh_oe:<26} D{row:<5} driver-built EBITA less bridge EBITA "
              f"(FY2025A) = {v:.1f}")
        print("          Expected -2.0: $310M total purchased-intangible amortisation "
              "is added back, but only")
        print("          $151M + $157M = $308M is identifiable by line. Disclosed in "
              "the workbook, not hidden.")
        if abs(v + 2.0) > 0.5:
            fails.append(("EBITA reconciliation gap", f"D{row}", v, -2.0))

    print()
    print("=" * 96)
    print("PART 2 — MODEL vs PUBLISHED WORK PRODUCTS")
    print("=" * 96)
    print(f"  {'Check':<58} {'Model':>13} {'Published':>12} {'Diff':>10}")
    print("  " + "-" * 94)
    for sheet in (sh_oe, sh_tv):
        w = wb[sheet]
        hdr = find_rows(w, "CHECKS AGAINST THE PUBLISHED")
        if not hdr:
            continue
        print(f"\n  --- {sheet} ---")
        for row in range(hdr[0], w.max_row + 1):
            lab = w.cell(row=row, column=1).value
            pub = w.cell(row=row, column=3).value
            if not isinstance(lab, str) or not isinstance(pub, (int, float)):
                continue
            mod = get(sheet, f"B{row}")
            if mod is None:
                continue
            diff = mod - pub
            tol = max(TOL_ABS, abs(pub) * TOL_REL)
            ok = abs(diff) <= tol
            flag = "OK " if ok else "FAIL"
            print(f"  [{flag}] {lab[:56]:<58} {mod:>13,.2f} {pub:>12,.2f} "
                  f"{diff:>10,.2f}")
            if not ok:
                fails.append((lab, f"{sheet}!B{row}", mod, pub))
            elif abs(diff) > 0.51:
                warns.append((lab, diff))

    print()
    print("=" * 96)
    print("PART 3 — HEADLINE OUTPUTS")
    print("=" * 96)
    w = wb[sh_tv]
    ps = find_rows(w, "INTRINSIC VALUE PER DILUTED SHARE")
    for row in ps:
        lab = w.cell(row=row, column=1).value.strip()
        cols = [c for c in "BCDEF"
                if w.cell(row=row,
                          column=openpyxl.utils.column_index_from_string(c)).value]
        vs = {c: get(sh_tv, f"{c}{row}") for c in cols}
        if len(cols) == 1:
            print(f"  {lab:<52} ${vs['B']:,.2f}")
        else:
            print(f"  {lab:<52} " + "  ".join(
                f"{n}=${vs[c]:,.2f}" for c, n in
                zip(cols, ["SEVERE", "BEAR", "BASE", "BULL", "WEIGHTED"])))
    oe = wb[sh_oe]
    for row in find_rows(oe, "= OWNERS' EARNINGS"):
        if oe.cell(row=row, column=2).value is None:
            continue
        vs = [get(sh_oe, f"{c}{row}") for c in "BCDEFGHIJ"]
        if any(v is None for v in vs):
            continue
        print(f"  Owners' earnings FY23A/24A/25A/TTM -> FY26E..FY30E ($M):")
        print("    " + "  ".join(f"{v:,.0f}" for v in vs))
        break

    print()
    print("=" * 96)
    if fails:
        print(f"RESULT: {len(fails)} FAILURE(S)")
        for f in fails:
            print("   ", f)
        return 1
    print(f"RESULT: ALL CHECKS PASS  ({len(warns)} residual(s) above $0.51 from "
          f"published rounding, all within tolerance)")
    for lab, d in warns:
        print(f"    residual {d:+.2f}  {lab[:70]}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
